# -*- coding: utf-8 -*-
"""
文件格式适配层(dt_files)
========================
xlsx / csv / json 三类文件的读写与值转换:

* 表头解析与行对齐(首行表头 ↔ 数据库字段名,sheet 名 / 一级键 ↔ 表名);
* 数值保真:数值一律经 dt_numeric 做 Decimal 定点转换,禁止 float 中转与科学计数法;
* Excel 超过双精度精度或带尾零的数值按文本写出,避免被 Excel 截断 / 舍入;
* CSV 支持分隔符、编码、多表目录模式与旧版分节文件(向后兼容)。
"""

from __future__ import annotations

import contextlib
import csv
import json
import os
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from math import isfinite

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

try:  # openpyxl 内部正则,用于剔除 Excel 不允许的控制字符
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
except Exception:  # pragma: no cover - 老版本 openpyxl 兜底
    ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")

from dt_config import CSV_SECTION_RE, EXCEL_MAX_CELL_CHARS, _close_quietly
from dt_numeric import (EXCEL_SIGNIFICANT_LIMIT, dumps_json, excel_cell_value,
                        format_number, parse_number, read_json_file)


# ---------------------------------------------------------------------------
# 值的类型转换(数据库 / 文件 → 各类输出)
# ---------------------------------------------------------------------------
def decode_bytes(value) -> str:
    """bytes → 文本(非法字节用替换符兜底,绝不因编码问题中断导出)。"""
    try:
        return value.decode("utf-8")
    except UnicodeDecodeError:
        return value.decode("utf-8", "replace")


def format_timedelta(value: timedelta) -> str:
    """timedelta → HH:MM:SS(MySQL TIME 列会返回 timedelta)。"""
    negative = value < timedelta(0)
    hours, remain = divmod(int(abs(value.total_seconds())), 3600)
    minutes, seconds = divmod(remain, 60)
    return f"{'-' if negative else ''}{hours:02d}:{minutes:02d}:{seconds:02d}"


def complex_to_json_text(value) -> str:
    """dict / list / tuple 等复合值(如 PostgreSQL jsonb 列、MongoDB 嵌套文档)→ JSON 文本。"""
    return json.dumps(value, ensure_ascii=False, default=str)


def sanitize_excel_text(text: str) -> str:
    """剔除 Excel 不允许的控制字符,并截断到单元格字符上限。"""
    return ILLEGAL_CHARACTERS_RE.sub("", text)[:EXCEL_MAX_CELL_CHARS]


def excel_number(value):
    """Excel 单元格数值兜底。

    Excel 内部是 IEEE-754 双精度:NaN / Infinity 无法表示,超过 15 位有效数字的
    整数与 Decimal(以及带尾零的小数)会被截断或舍入,因此这些值改以文本写出,
    保证与输入字面量完全一致。
    """
    if isinstance(value, Decimal):
        return excel_cell_value(value)
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return str(value) if len(str(abs(value))) > EXCEL_SIGNIFICANT_LIMIT else value
    if isinstance(value, float) and not isfinite(value):
        return str(value)
    return value


def to_excel_value(value):
    """任意来源值 → 可安全写入 Excel 单元格的值。"""
    if value is None:
        return None
    if isinstance(value, str):
        return sanitize_excel_text(value)
    if isinstance(value, datetime):
        # Excel 不支持带时区的日期,去掉时区信息(保留驱动返回的墙上时间)
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (date, Decimal, int, float, bool)):
        return excel_number(value)
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        return format_timedelta(value)
    if isinstance(value, (dict, list, tuple)):
        return sanitize_excel_text(complex_to_json_text(value))
    if isinstance(value, (bytes, bytearray)):
        return sanitize_excel_text(decode_bytes(bytes(value)))
    return sanitize_excel_text(str(value))


def to_json_value(value):
    """数据库 / 文件返回值 → JSON 兼容值(时间统一为字符串)。

    ``Decimal`` 原样保留,由 :func:`dt_numeric.dumps_json` 以定点数字写出,
    避免 float 中转造成精度丢失或科学计数法。
    """
    if value is None:
        return None
    if isinstance(value, dict):                   # json/jsonb / 嵌套文档:递归转换
        return {str(key): to_json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [to_json_value(item) for item in value]
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        return format_timedelta(value)
    if isinstance(value, (bytes, bytearray)):
        return decode_bytes(bytes(value))
    if isinstance(value, (int, float, str, bool, Decimal)):
        return value
    return str(value)                             # ObjectId / UUID / 其他类型


def to_csv_value(value):
    """任意来源值 → CSV 单元格文本(数值一律定点输出,禁止科学计数法)。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        return format_timedelta(value)
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return format_number(value)
    if isinstance(value, (dict, list, tuple)):   # json/jsonb / 嵌套文档
        return complex_to_json_text(value)
    if isinstance(value, (bytes, bytearray)):
        return decode_bytes(bytes(value))
    return value


def to_db_value(value):
    """任意来源值 → 可交给数据库驱动的值(Decimal 直接交给驱动,不经 float)。"""
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (dict, list, tuple)):   # json/jsonb 列 → JSON 文本入库
        return complex_to_json_text(value)
    if isinstance(value, (bytes, bytearray)):
        return bytes(value)
    return value  # datetime / date / timedelta / Decimal / bool / int / float


# ---------------------------------------------------------------------------
# 行 / 表头处理
# ---------------------------------------------------------------------------
def is_blank_cell(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_blank_row(row) -> bool:
    return all(map(is_blank_cell, row))


def normalize_row(row, ncols: int) -> list:
    """把行对齐到表头列数:短了补 None,长了截断。"""
    values = list(row)
    return values[:ncols] if len(values) > ncols else values + [None] * (ncols - len(values))


def parse_header(raw_row) -> list:
    """解析表头:去掉尾部空列,空列自动命名,重名自动加后缀(保证唯一)。"""
    cells = list(raw_row)
    end = len(cells)
    while end and is_blank_cell(cells[end - 1]):
        end -= 1
    if not end:
        raise ValueError("数据表头为空,无法转换")

    header, seen = [], set()
    for index, cell in enumerate(cells[:end]):
        base = (str(cell).strip() if cell is not None else "") or f"column_{index + 1}"
        name, suffix = base, 1
        while name in seen:                  # 自动后缀再次撞名时继续递增
            suffix += 1
            name = f"{base}_{suffix}"
        seen.add(name)
        header.append(name)
    return header


def display_width(text) -> int:
    """按显示宽度估算列宽(中文等全角字符算 2)。"""
    return sum(2 if ord(char) > 0x7F else 1 for char in str(text))


def safe_sheet_title(name) -> str:
    """工作表名:去掉非法字符,长度 <= 31。"""
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", str(name)).strip()
    return cleaned[:31] or "Sheet1"


def safe_filename(name) -> str:
    """去掉文件名中 Windows 不允许的字符。"""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', "_", str(name)).strip()
    return cleaned or "export"


def filter_columns(header, rows_iter, columns):
    """按勾选字段过滤行流,返回 (新表头, 行生成器)。"""
    indexes = [header.index(column) for column in columns if column in header]
    new_header = [header[index] for index in indexes]

    def gen():
        for row in rows_iter:
            yield [row[index] for index in indexes]

    return new_header, gen()


# ---------------------------------------------------------------------------
# Excel 读写(openpyxl 流式模式)
# ---------------------------------------------------------------------------
def write_rows_to_worksheet(ws, header, row_iterator, log=None) -> int:
    """把数据行逐行写入工作表,顺便自动计算列宽。"""
    ws.append([to_excel_value(name) for name in header])
    widths = {index: min(display_width(name) + 2, 60)
              for index, name in enumerate(header, start=1)}
    total = truncated = 0
    for row in row_iterator:
        raw_values = list(row)
        values = [to_excel_value(value) for value in raw_values]
        ws.append(values)
        total += 1
        for index, value in enumerate(values, start=1):
            if value is None:
                continue
            width = min(display_width(value) + 2, 60)
            if width > widths.get(index, 0):
                widths[index] = width
        truncated += sum(1 for raw in raw_values
                         if isinstance(raw, str) and len(raw) > EXCEL_MAX_CELL_CHARS)
        if log and total % 20000 == 0:
            log(f"  已写出 {total:,} 行 ...")
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = width
    if truncated and log:
        log(f"  ⚠ 有 {truncated} 个单元格超过 Excel 上限 {EXCEL_MAX_CELL_CHARS} 字符,已截断")
    return total


def write_excel_file(path, sheet, columns, row_iterator, mode="replace", log=None) -> int:
    """写出 Excel 文件。

    mode=replace : 用流式 write_only 模式重建整个工作簿(大数据量首选);
    mode=merge   : 保留原工作簿其他 Sheet,覆盖同名 Sheet(无同名则新增);
                   目标文件不存在时自动退化为 replace。
    """
    if mode == "merge" and not os.path.exists(path):
        mode = "replace"
    if mode == "replace":
        wb = Workbook(write_only=True)
        try:
            ws = wb.create_sheet(title=sheet)
            total = write_rows_to_worksheet(ws, columns, row_iterator, log)
            wb.save(path)
        finally:
            _close_quietly(wb)
        return total

    wb = load_workbook(path)          # merge:只动同名 Sheet
    try:
        if sheet in wb.sheetnames:
            del wb[sheet]
        ws = wb.create_sheet(title=sheet)
        total = write_rows_to_worksheet(ws, columns, row_iterator, log)
        wb.save(path)
    finally:
        _close_quietly(wb)
    return total


def write_excel_multi_file(path, tables, mode="replace", log=None) -> int:
    """多表写出 Excel:tables 为可迭代的 (sheet_name, columns, row_iter)。

    每个表一个 Sheet;mode=merge 时保留原工作簿其他 Sheet,覆盖同名 Sheet;
    目标文件不存在时自动退化为 replace。
    """
    if mode != "replace" and not os.path.exists(path):
        mode = "replace"
    if mode != "replace":
        return _write_excel_multi_merge(path, tables, log)
    wb = Workbook(write_only=True)
    try:
        total = 0
        for sheet, columns, rows in tables:
            total += write_rows_to_worksheet(wb.create_sheet(title=sheet), columns, rows, log)
        wb.save(path)
    finally:
        _close_quietly(wb)
    return total


def _write_excel_multi_merge(path, tables, log) -> int:
    wb = load_workbook(path)
    try:
        total = 0
        for sheet, columns, rows in tables:
            if sheet in wb.sheetnames:
                del wb[sheet]
            total += write_rows_to_worksheet(wb.create_sheet(title=sheet), columns, rows, log)
        wb.save(path)
    finally:
        _close_quietly(wb)
    return total


def excel_cell_to_value(value):
    """Excel 单元格 → 保真值:浮点通过十进制字面量转换,不让二进制误差进入后续环节。"""
    return Decimal(str(value)) if isinstance(value, float) else value


def load_excel_stream(path, sheet):
    """读取 xlsx 工作表,返回 (表头, 行生成器)。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"工作簿中不存在工作表 {sheet!r},可用: {', '.join(wb.sheetnames)}")
        rows = wb[sheet].iter_rows(values_only=True)
        if (raw_header := next(rows, None)) is None:
            raise ValueError("工作表为空,无法读取")
        header = parse_header(raw_header)
    except Exception:
        _close_quietly(wb)
        raise

    def gen():
        try:
            for row in rows:
                if not is_blank_row(row):
                    yield normalize_row([excel_cell_to_value(cell) for cell in row], len(header))
        finally:
            _close_quietly(wb)

    return header, gen()


def workbook_sheet_headers(path) -> dict:
    """读取工作簿每个 Sheet 的表头;表头为空的工作表返回 [](不让整份文件加载失败)。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        headers = {}
        for name in wb.sheetnames:
            first = next(wb[name].iter_rows(values_only=True, max_row=1), None)
            try:
                headers[name] = parse_header(first) if first else []
            except ValueError:
                headers[name] = []
        return headers
    finally:
        _close_quietly(wb)


# ---------------------------------------------------------------------------
# JSON 读写(行记录式结构 + 旧版列数组兼容)
# ---------------------------------------------------------------------------
def build_json_document(top_key, columns, rows) -> dict:
    """构造 {表名: {HeaderFields: [字段...], Data: {RowN: [值...]}}} 的行记录式 JSON 结构。"""
    header_fields = [str(column) for column in columns]
    data = {f"Row{index}": [to_json_value(value) for value in normalize_row(row, len(columns))]
            for index, row in enumerate(rows, start=1)}
    return {top_key: {"HeaderFields": header_fields, "Data": data}}


def _json_merge_base(path, payload: dict) -> dict:
    """merge 模式:读出原文件的一级键并覆盖同名键(读失败则以新内容为准)。"""
    with contextlib.suppress(Exception):
        old = read_json_file(path)
        if isinstance(old, dict):
            old.update(payload)
            return old
    return payload


def build_json_documents(top_key, columns, rows) -> dict:
    """构造 {集合名: [{字段: 值, ...}, ...]} —— MongoDB 文档的直接映射结构。

    MongoDB 的 collection 与 JSON 的对象数组天然对应,嵌套子文档 / 数组保持原结构;
    MySQL / PostgreSQL 仍用 HeaderFields + Data 的行记录式结构(见 build_json_document)。
    """
    header_fields = [str(column) for column in columns]
    documents = []
    for row in rows:
        values = normalize_row(row, len(columns))
        documents.append({name: to_json_value(value)
                          for name, value in zip(header_fields, values)})
    return {top_key: documents}


def write_json_file(path, top_key, columns, rows, mode="replace", log=None,
                    documents=False) -> int:
    """写出 JSON 文件。mode=merge 时保留原文件其他一级键,仅覆盖同名键。

    documents=True(MongoDB 来源)写出 ``{键: [文档, ...]}`` 直接映射结构,
    否则写行记录式 ``{键: {HeaderFields: [...], Data: {RowN: [...]}}}``。
    """
    builder = build_json_documents if documents else build_json_document
    document = builder(top_key, columns, rows)
    if mode == "merge" and os.path.exists(path):
        document = _json_merge_base(path, document)
    with open(path, "w", encoding="utf-8") as fp:
        fp.write(dumps_json(document))               # 定点数字,禁止科学计数法
    return len(document[top_key]) if documents else len(document[top_key]["Data"])


def write_json_multi_file(path, tables, mode="replace", log=None) -> int:
    """多表写出 JSON:tables 为可迭代的 (top_key, columns, rows[, documents])。

    每个表一个一级键;documents 为真时该键写成文档数组(MongoDB 直接映射),
    否则写行记录式(HeaderFields/Data);mode=merge 时保留原文件其他一级键。
    """
    payload, total = {}, 0
    for entry in tables:
        key, columns, rows = entry[0], entry[1], entry[2]
        documents = bool(len(entry) > 3 and entry[3])
        builder = build_json_documents if documents else build_json_document
        sub = builder(key, columns, rows)[key]
        payload[key] = sub
        total += len(sub) if documents else len(sub["Data"])
    if mode == "merge" and os.path.exists(path):
        payload = _json_merge_base(path, payload)
    with open(path, "w", encoding="utf-8") as fp:
        fp.write(dumps_json(payload))
    return total


def load_json_stream(path, top_key):
    """读取 JSON 表结构,返回 (表头, 行生成器)。

    支持三种结构:
    * 行记录式(关系库默认输出):{表名: {HeaderFields: [字段...], Data: {RowN: [值...]}}};
    * 文档数组式(MongoDB 直接映射输出):{表名: [{字段: 值, ...}, ...]};
    * 旧版列数组:{表名: {字段名: [值...]}} 仍可读取(向后兼容)。

    小数以 ``Decimal`` 读入(``parse_float=Decimal``),长浮点数不会退化为 float。
    """
    doc = read_json_file(path)
    if not isinstance(doc, dict):
        raise ValueError("JSON 顶层必须是对象")
    if top_key not in doc:
        raise ValueError(f"JSON 中不存在一级键 {top_key!r},可用: {list(doc.keys())}")
    table = doc[top_key]

    if isinstance(table, list):
        # MongoDB 直接映射结构:{表名: [{字段: 值, ...}, ...]}
        rows = [row for row in table if isinstance(row, dict)]
        if len(rows) != len(table):
            raise ValueError(f"一级键 {top_key!r} 的数组元素必须是对象")
        columns = []
        for row in rows:
            for name in row:
                if str(name) not in columns:
                    columns.append(str(name))

        def gen_documents():
            for row in rows:
                yield [row.get(name) for name in columns]

        return columns, gen_documents()

    if not isinstance(table, dict):
        raise ValueError(f"一级键 {top_key!r} 的值必须是对象或对象数组")

    if isinstance(table.get("HeaderFields"), list) and isinstance(table.get("Data"), dict):
        # 行记录式结构:HeaderFields + Data.RowN
        columns = [str(column) for column in table["HeaderFields"]]
        data = table["Data"]
        keys = sorted(data.keys(), key=_row_key_rank)

        def gen():
            for key in keys:
                row = data[key]
                if not isinstance(row, list):
                    raise ValueError(f"Data 中 {key!r} 的值必须是数组")
                yield normalize_row(row, len(columns))

        return columns, gen()

    # 旧版列数组结构:字段名 -> 值数组
    columns = list(table.keys())
    arrays = []
    for column in columns:
        array = table[column]
        if not isinstance(array, list):
            raise ValueError(f"字段 {column!r} 的值必须是数组")
        arrays.append(array)
    row_count = max(map(len, arrays), default=0)

    def gen_legacy():
        for index in range(row_count):
            yield [array[index] if index < len(array) else None for array in arrays]

    return columns, gen_legacy()


def _row_key_rank(key) -> tuple:
    """行键排序权重:Row1/Row2/... 按数字升序,无数字后缀的键按文本序排后。"""
    match = re.search(r"\d+$", str(key))
    return (0, int(match.group())) if match else (1, str(key))


def json_top_keys(path) -> list:
    """读取 JSON 文件的一级键。"""
    doc = read_json_file(path)
    if not isinstance(doc, dict):
        raise ValueError("JSON 顶层必须是对象")
    return [str(key) for key in doc]


# ---------------------------------------------------------------------------
# CSV 读写(含多表目录模式与旧版分节文件的向后兼容)
# ---------------------------------------------------------------------------
def write_csv_file(path, columns, row_iterator, delimiter=",", encoding="utf-8-sig",
                   mode="replace", log=None) -> int:
    """写出 CSV 文件。

    mode=replace : 重建文件(含表头);
    mode=merge   : 向已有文件追加数据行(不重复写表头)。
    """
    append = mode == "merge" and os.path.exists(path)
    total = 0
    with open(path, "a" if append else "w", newline="", encoding=encoding) as fp:
        writer = csv.writer(fp, delimiter=delimiter, lineterminator="\n")
        if not append:
            writer.writerow([to_csv_value(name) for name in columns])
        for row in row_iterator:
            writer.writerow([to_csv_value(value)
                             for value in normalize_row(row, len(columns))])
            total += 1
            if log and total % 50000 == 0:
                log(f"  已写出 {total:,} 行 ...")
    return total


def read_csv_header(path, delimiter, encoding) -> list:
    with open(path, "r", newline="", encoding=encoding) as fp:
        reader = csv.reader(fp, delimiter=delimiter)
        if (raw := next(reader, None)) is None:
            raise ValueError("CSV 文件为空,无法读取")
    return parse_header(raw)


def parse_csv_cell(value):
    """CSV 单元格文本 → 自动推断 int / Decimal / date / datetime,失败保持字符串。

    * 长浮点数一律解析为 ``Decimal``(绝不产生 float),科学计数法输入按定点数值处理;
    * ``007`` 这类补零字面量按编号保留文本,避免丢失前导零;
    * 整数使用 Python 任意精度 ``int``,超大整数不会失真。
    """
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return None
    if (number := parse_number(text)) is not None:
        return number
    try:
        return datetime.strptime(text, "%Y-%m-%d %H:%M:%S")
    except ValueError:
        pass
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    return text


def iter_csv_rows(path, delimiter, encoding, header_len, parse=False):
    """逐行生成 CSV 数据(跳过表头与空行),parse=True 时按类型解析单元格。"""
    with open(path, "r", newline="", encoding=encoding) as fp:
        reader = csv.reader(fp, delimiter=delimiter)
        next(reader, None)                        # 跳过表头
        for row in reader:
            if is_blank_row(row):
                continue
            norm = normalize_row(row, header_len)
            yield [parse_csv_cell(cell) for cell in norm] if parse else norm


def scan_csv_sections(path, delimiter=",", encoding="utf-8-sig") -> list:
    """流式扫描旧版分节 CSV 的分节名(不把整个文件读进内存)。"""
    with open(path, "r", newline="", encoding=encoding) as fp:
        return [match.group(1)
                for row in csv.reader(fp, delimiter=delimiter)
                if len(row) == 1 and (match := CSV_SECTION_RE.match(row[0].strip()))]


def read_csv_sections(path, delimiter, encoding):
    """解析旧版分节 CSV,返回 {分节名: (表头, 数据行列表)};普通 CSV 返回 None。"""
    with open(path, "r", newline="", encoding=encoding) as fp:
        rows_all = list(csv.reader(fp, delimiter=delimiter))
    if not any(len(row) == 1 and CSV_SECTION_RE.match(row[0].strip()) for row in rows_all):
        return None
    sections, current, header, rows = {}, None, None, []
    for row in rows_all:
        match = CSV_SECTION_RE.match(row[0].strip()) if len(row) == 1 else None
        if match:
            if current is not None:
                sections[current] = (header or ["column_1"], rows)
            current, header, rows = match.group(1), None, []
        elif current is not None:
            if header is None:
                header = parse_header(row)
            elif not is_blank_row(row):
                rows.append(row)
    if current is not None:
        sections[current] = (header or ["column_1"], rows)
    return sections


def load_csv_section(path, section, delimiter, encoding, parse=False):
    """读取分节 CSV 中的某一节,返回 (表头, 行生成器)。"""
    sections = read_csv_sections(path, delimiter, encoding)
    if not sections or section not in sections:
        raise ValueError(f"分节 CSV 中不存在分节 {section!r},可用: {', '.join(sections or [])}")
    header, rows = sections[section]

    def gen():
        for row in rows:
            norm = normalize_row(row, len(header))
            yield [parse_csv_cell(cell) for cell in norm] if parse else norm

    return header, gen()
