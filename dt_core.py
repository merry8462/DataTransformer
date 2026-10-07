# -*- coding: utf-8 -*-
"""
核心转换引擎与统一门面(dt_core)
================================
本模块是 DataTransformer 的核心转换引擎:
* 任务规格:SourceSpec / TargetSpec / FieldMapping / RunResult;
* Reader / Writer 统一接口(open_reader / convert):数据库与文件互转共用一条管道;
* 数值保真回读校验(写入后比对,失败即阻断并输出差异报告);
* SSH 隧道与无界面自检。

同时作为门面统一再导出 dt_config / dt_files / dt_sql / dt_mongo / dt_numeric 的
公开名称,界面层(dt_cli.py、data_transformer.py)只需 `import dt_core`。
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import socket
import threading
from dataclasses import dataclass, field, replace
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    import pymysql
except Exception:  # pragma: no cover - 未安装驱动时仍可使用文件互转
    pymysql = None

try:
    import psycopg2
except Exception:  # pragma: no cover
    psycopg2 = None

try:
    import paramiko
except Exception:  # pragma: no cover - 未安装 paramiko 时仅无法使用 SSH 隧道
    paramiko = None

from dt_config import (
    APP_TITLE, VERSION, SAMPLE_SIZE, DEFAULT_BATCH, EXCEL_MAX_CELL_CHARS,
    SSH_DEFAULT_PORT, INPUT_LABELS, OUTPUT_LABELS, MODE_APPEND, MODE_CREATE,
    MODE_REPLACE, MODE_LABELS, MODE_KEYS, MODES, DB_KINDS, FILE_KINDS,
    DELIM_LABELS, ENC_LABELS, IDENT_RE, CSV_SECTION_RE, log_noop,
    config_file_path, to_int, DBConfig, SSHConfig, Settings, _close_quietly,
)
from dt_files import (
    decode_bytes, format_timedelta, complex_to_json_text, sanitize_excel_text,
    excel_number, to_excel_value, to_json_value, to_csv_value, to_db_value,
    is_blank_cell, is_blank_row, normalize_row, parse_header, display_width,
    safe_sheet_title, safe_filename, write_rows_to_worksheet, write_excel_file,
    write_excel_multi_file, excel_cell_to_value, load_excel_stream,
    workbook_sheet_headers, build_json_document, write_json_file, write_json_multi_file,
    load_json_stream, json_top_keys, write_csv_file, read_csv_header, parse_csv_cell,
    iter_csv_rows, scan_csv_sections, read_csv_sections, load_csv_section, filter_columns,
)
from dt_sql import (
    DBAdapter, MySQLAdapter, PostgresAdapter, ADAPTERS, describe_db_error,
    close_cursor, load_sql_stream, take_sample, write_sql_table, ColumnStats,
    decide_type, collect_stats, build_create_table,
)
from dt_numeric import (
    DEFAULT_SAMPLE_LIMIT, FidelityError, FidelityReport, NumericRecorder,
)
from dt_mongo import (
    MongoConfig, load_mongo_stream, pymongo_available, write_mongo_collection,
)

# 门面公开接口:界面层(dt_cli.py / data_transformer.py)与测试只需 `import dt_core`
__all__ = [
    # 常量与配置(dt_config)
    "APP_TITLE", "VERSION", "SAMPLE_SIZE", "DEFAULT_BATCH", "EXCEL_MAX_CELL_CHARS",
    "SSH_DEFAULT_PORT", "INPUT_LABELS", "OUTPUT_LABELS", "MODE_APPEND", "MODE_CREATE",
    "MODE_REPLACE", "MODE_LABELS", "MODE_KEYS", "MODES", "DB_KINDS", "FILE_KINDS",
    "DELIM_LABELS", "ENC_LABELS", "IDENT_RE", "CSV_SECTION_RE", "log_noop",
    "config_file_path", "to_int", "DBConfig", "SSHConfig", "Settings",
    # 文件格式适配层(dt_files)
    "decode_bytes", "format_timedelta", "complex_to_json_text", "sanitize_excel_text",
    "excel_number", "to_excel_value", "to_json_value", "to_csv_value", "to_db_value",
    "is_blank_cell", "is_blank_row", "normalize_row", "parse_header", "display_width",
    "safe_sheet_title", "safe_filename", "write_rows_to_worksheet", "write_excel_file",
    "write_excel_multi_file", "excel_cell_to_value", "load_excel_stream",
    "workbook_sheet_headers", "build_json_document", "write_json_file",
    "write_json_multi_file", "load_json_stream", "json_top_keys", "write_csv_file",
    "read_csv_header", "parse_csv_cell", "iter_csv_rows", "scan_csv_sections",
    "read_csv_sections", "load_csv_section", "filter_columns",
    # 数据库适配层(dt_sql)与 MongoDB(dt_mongo)
    "DBAdapter", "MySQLAdapter", "PostgresAdapter", "ADAPTERS", "describe_db_error",
    "close_cursor", "load_sql_stream", "take_sample", "write_sql_table", "ColumnStats",
    "decide_type", "collect_stats", "build_create_table", "MongoConfig",
    "load_mongo_stream", "write_mongo_collection", "pymongo_available",
    # 数值保真层(dt_numeric)
    "FidelityError", "FidelityReport", "NumericRecorder", "DEFAULT_SAMPLE_LIMIT",
    # 转换引擎(dt_core)
    "paramiko_available", "describe_ssh_error", "SSHTunnel", "config_for",
    "connect_with_ssh", "FieldMapping", "SourceSpec", "TargetSpec", "RunResult",
    "csv_output_path", "csv_output_paths", "open_reader", "apply_field_mapping",
    "recording_rows", "readback_source", "verify_readback", "convert", "run_self_test",
]





def paramiko_available() -> bool:
    return paramiko is not None


def describe_ssh_error(exc, cfg=None) -> str:
    """把 SSH 异常翻译成便于用户定位的中文提示。"""
    raw = str(exc)
    low = raw.lower()
    view = cfg.as_dict() if isinstance(cfg, SSHConfig) else dict(cfg or {})
    host = view.get("host") or "?"
    port = view.get("port") or SSH_DEFAULT_PORT
    user = view.get("user") or "?"
    if isinstance(exc, (ValueError, TypeError)):
        return raw
    if paramiko is not None and isinstance(exc, paramiko.AuthenticationException):
        return (f"SSH 认证失败:{user}@{host}\n"
                f"请检查【SSH 用户名】【SSH 密码】或私钥文件。\n原始错误: {raw}")
    if any(key in low for key in ("authentication failed", "auth fail", "permission denied")):
        return (f"SSH 认证失败:{user}@{host}\n"
                f"请检查【SSH 用户名】【SSH 密码】或私钥文件。\n原始错误: {raw}")
    if any(key in low for key in ("timed out", "timeout")):
        return f"SSH 连接 {host}:{port} 超时。\n请检查主机、端口、网络与防火墙。\n原始错误: {raw}"
    if "refused" in low:
        return f"SSH 连接被拒绝:{host}:{port}。\n请确认目标机 sshd 已启动且端口正确。\n原始错误: {raw}"
    if any(key in low for key in ("name or service not known", "getaddrinfo", "gaierror")):
        return f"SSH 主机名解析失败:{host!r}。\n请检查【SSH 主机】。\n原始错误: {raw}"
    if paramiko is not None and isinstance(exc, paramiko.SSHException):
        return f"SSH 通道异常:{raw}"
    return raw


def _pump(src, dst) -> None:
    """把 src 收到的字节转发给 dst,任一端结束即关闭双向连接。"""
    try:
        while data := src.recv(32768):
            dst.sendall(data)
    except Exception:
        pass
    finally:
        _close_quietly(src)
        _close_quietly(dst)


class SSHTunnel:
    """SSH 本地端口转发:127.0.0.1:<local_port> → 远程网络中的 <remote_host>:<remote_port>。

    转发在独立线程中双向搬运字节,因此 MySQL / PostgreSQL 适配器无需任何改动,
    只要把数据库主机指向 127.0.0.1、端口指向 local_port 即可穿透 SSH 操作远端数据库。
    """

    BIND_HOST = "127.0.0.1"

    def __init__(self, cfg: SSHConfig, remote_host: str = "127.0.0.1", remote_port: int = 0):
        self.cfg = cfg
        self.remote_host = (remote_host or "127.0.0.1").strip()
        self.remote_port = to_int(remote_port, 0)
        self.local_port = 0
        self._client = None
        self._transport = None
        self._listener = None
        self._stop = threading.Event()

    # -- 生命周期 ----------------------------------------------------------
    def start(self) -> "SSHTunnel":
        if paramiko is None:
            raise RuntimeError("未安装 paramiko,无法使用 SSH 隧道。请先执行: pip install paramiko")
        cfg = self.cfg.validated()
        if not self.remote_port:
            raise ValueError("未指定远程数据库端口,无法建立 SSH 隧道")
        client = paramiko.SSHClient()
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        options = {
            "hostname": cfg.host,
            "port": cfg.port,
            "username": cfg.user,
            "timeout": cfg.timeout,
            "banner_timeout": cfg.timeout,
            "auth_timeout": cfg.timeout,
            "allow_agent": True,
            "look_for_keys": not any((cfg.password, cfg.key_file)),
        }
        if cfg.password:
            options["password"] = cfg.password
        if cfg.key_file:
            options["key_filename"] = os.path.expanduser(cfg.key_file)
            if cfg.passphrase:
                options["passphrase"] = cfg.passphrase
        try:
            client.connect(**options)
        except Exception:
            _close_quietly(client)
            raise
        transport = client.get_transport()
        if transport is None or not transport.is_active():
            _close_quietly(client)
            raise RuntimeError(f"SSH 通道建立失败:{cfg.host}:{cfg.port}")
        transport.set_keepalive(30)
        listener = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        listener.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        listener.bind((self.BIND_HOST, 0))
        listener.listen(16)
        listener.settimeout(0.5)
        self._client, self._transport, self._listener = client, transport, listener
        self.local_port = listener.getsockname()[1]
        self._stop.clear()
        threading.Thread(target=self._accept_loop, name="dt-ssh-tunnel",
                         daemon=True).start()
        return self

    def stop(self) -> None:
        self._stop.set()
        for handle in (self._listener, self._transport, self._client):
            _close_quietly(handle)
        self._listener = self._transport = self._client = None
        self.local_port = 0

    @property
    def active(self) -> bool:
        transport = self._transport
        return bool(transport is not None and self._listener is not None
                    and transport.is_active())

    def describe(self) -> str:
        cfg = self.cfg
        return (f"{self.BIND_HOST}:{self.local_port} → SSH {cfg.user}@{cfg.host}:{cfg.port}"
                f" → {self.remote_host}:{self.remote_port}")

    # -- 内部实现 ----------------------------------------------------------
    def _accept_loop(self) -> None:
        listener = self._listener
        while not self._stop.is_set() and listener is not None:
            try:
                conn, _ = listener.accept()
            except socket.timeout:
                continue
            except OSError:
                break
            threading.Thread(target=self._forward, args=(conn,),
                             name="dt-ssh-conn", daemon=True).start()

    def _forward(self, conn) -> None:
        """把本地 TCP 连接接到 SSH 的 direct-tcpip 通道上。"""
        try:
            channel = self._transport.open_channel(
                "direct-tcpip", (self.remote_host, self.remote_port), conn.getpeername())
        except Exception:
            _close_quietly(conn)
            return
        threading.Thread(target=_pump, args=(conn, channel),
                         name="dt-ssh-up", daemon=True).start()
        _pump(channel, conn)


def config_for(adapter_name: str, values=None):
    """按适配器类型构造连接配置(MySQL / PostgreSQL → DBConfig;MongoDB → MongoConfig)。"""
    source = values.as_dict() if hasattr(values, "as_dict") else (values or {})
    return (MongoConfig.from_mapping(source) if adapter_name == "MongoDB"
            else DBConfig.from_mapping(source))


def connect_with_ssh(db_cfg: DBConfig, ssh_cfg=None, adapter=None, log=log_noop):
    """按需建立 SSH 隧道并连接数据库,返回 (连接, 隧道或 None)。

    开启 SSH 时 db_cfg.host/port 表示【远程服务器上】数据库的地址(通常 127.0.0.1),
    本地连接则改走隧道端口。MySQL / PostgreSQL / MongoDB 均适用。
    """
    tunnel = None
    if ssh_cfg is not None and ssh_cfg.enabled:
        tunnel = SSHTunnel(ssh_cfg, remote_host=db_cfg.host or "127.0.0.1",
                           remote_port=db_cfg.resolved_port(adapter)).start()
        log(f"SSH 隧道已建立: {tunnel.describe()}")
        db_cfg = replace(db_cfg, host=SSHTunnel.BIND_HOST, port=tunnel.local_port)
    try:
        return adapter.connect(db_cfg), tunnel
    except Exception:
        if tunnel is not None:
            tunnel.stop()
        raise


@dataclass
class FieldMapping:
    """字段映射:源字段名 → 目标字段名(默认同名,可手动覆盖)。

    对应需求中的映射约定:Excel 首行表头 / CSV 表头 / JSON 字段名 ↔ 数据库字段名。
    """

    pairs: dict = field(default_factory=dict)

    def target_of(self, source: str) -> str:
        return self.pairs.get(source, source)

    def rename_header(self, header) -> list:
        return [self.target_of(str(name)) for name in header]

    def describe(self) -> str:
        return "、".join(f"{src}→{dst}" for src, dst in self.pairs.items()) or "(同名映射)"

    @classmethod
    def from_pairs(cls, pairs=None) -> "FieldMapping":
        return cls({str(k): str(v) for k, v in (pairs or {}).items() if str(v).strip()})


@dataclass
class SourceSpec:
    """一个输入源:一张数据库表 / 集合 / Excel Sheet / JSON 一级键 / CSV 分节。"""

    kind: str                     # SQL | Mongo | Excel | JSON | CSV
    table: str = ""               # 表名 / 集合名 / Sheet 名 / 一级键 / 分节名
    columns: list = field(default_factory=list)   # 勾选字段(空 = 全部)
    path: str = ""
    delimiter: str = ","
    encoding: str = "utf-8-sig"
    parse: bool = False           # CSV 文本是否按类型解析(导入数据库时开启)
    mode: str = ""                # 覆盖 TargetSpec.mode("skip" = 跳过该源)
    rows: int = 0                 # 执行后回填
    status: str = "pending"       # pending | done | skipped

    def label(self) -> str:
        return self.table or (Path(self.path).name if self.path else self.kind)


@dataclass
class TargetSpec:
    """输出目标(数据库表 / 集合 / Excel Sheet / JSON 一级键 / CSV 文件或目录)。"""

    kind: str                     # SQL | Mongo | Excel | JSON | CSV
    mode: str = "replace"         # 数据库: replace/create_if_missing/append;文件: replace/merge
    path: str = ""                # 文件输出路径,或 CSV 目录模式下的目录
    table: str = ""               # 目标表名 / 集合名
    sheet: str = ""               # 目标 Sheet 名(Excel)
    key: str = ""                 # 目标一级键(JSON)
    directory: bool = False       # CSV:目录模式(每个源一个 {表名}.csv)
    delimiter: str = ","
    encoding: str = "utf-8-sig"
    batch: int = DEFAULT_BATCH
    fields: FieldMapping = field(default_factory=FieldMapping)   # 字段映射
    names: dict = field(default_factory=dict)                    # 源名 → 目标名(表/集合/Sheet)
    verify: bool = True           # 写入后回读做数值保真校验
    allow_mismatch: bool = False  # 校验未通过时是否继续(默认阻断)

    def target_name(self, source_name: str) -> str:
        return self.names.get(source_name, source_name)


@dataclass
class RunResult:
    """转换结果汇总(含数值保真校验报告)。"""

    rows: int = 0
    sources: int = 0
    skipped: int = 0
    outputs: list = field(default_factory=list)
    fidelity: FidelityReport = field(default_factory=FidelityReport)

    def summary(self) -> str:
        parts = [f"共 {self.rows:,} 行", f"{self.sources} 个数据源"]
        if self.skipped:
            parts.append(f"{self.skipped} 个已跳过")
        if self.outputs:
            parts.append("输出: " + "、".join(map(str, self.outputs)))
        if self.fidelity.checked_values:
            parts.append("数值校验" + ("通过" if self.fidelity.ok else "未通过"))
        return ",".join(parts)


def csv_output_path(dir_path, src: "SourceSpec", index: int = 1, target=None) -> str:
    """目录模式下单个输入源对应的 .csv 文件路径(文件名遵循 names 映射)。"""
    name = target.target_name(src.table) if target is not None else src.table
    return os.path.join(dir_path, f"{safe_filename(name or f'table{index}')}.csv")


def csv_output_paths(dir_path, sources, target=None) -> list:
    """CSV 目录模式下每个输入源对应的文件路径(用于写前冲突检测)。"""
    return [csv_output_path(dir_path, src, index, target)
            for index, src in enumerate(sources, start=1)]


def open_reader(src: SourceSpec, conn=None, adapter=None):
    """打开输入源,返回 (表头, 行生成器, 关闭函数)。

    这是转换引擎的 **Reader 统一接口**:数据库(SQL / MongoDB)与
    xlsx / csv / json 都通过它暴露相同的 ``(header, rows, close)`` 三元组。
    """
    if src.kind == "Mongo":
        if conn is None or adapter is None:
            raise ValueError("MongoDB 输入需要先连接数据库")
        header, rows, closer = load_mongo_stream(conn, adapter, src.table, src.columns)
        return header, rows, closer
    if src.kind == "SQL":
        if conn is None or adapter is None:
            raise ValueError("SQL 输入需要先连接数据库")
        header, rows, cur = load_sql_stream(conn, adapter, src.table, src.columns)
        return header, rows, lambda: close_cursor(conn, cur)

    if src.kind == "Excel":
        header, rows = load_excel_stream(src.path, src.table)
    elif src.kind == "JSON":
        header, rows = load_json_stream(src.path, src.table)
    elif src.kind == "CSV":
        header, rows = _open_csv(src)
    else:
        raise ValueError(f"不支持的输入类型: {src.kind!r}")

    if selected := [name for name in src.columns if name in header]:
        if selected != header:
            header, rows = filter_columns(header, rows, selected)
    return header, rows, log_noop


def _open_csv(src: SourceSpec):
    """CSV 输入:``table`` 命中分节名时读取该分节,否则整份文件作为一张表。"""
    section = src.table.strip()
    if not section or section.startswith("("):
        return _read_plain_csv(src)
    sections = scan_csv_sections(src.path, src.delimiter, src.encoding)
    if section in sections:
        return load_csv_section(src.path, section, src.delimiter, src.encoding, parse=src.parse)
    if sections:
        raise ValueError(f"分节 CSV 中不存在分节 {section!r},可用: {', '.join(sections)}")
    return _read_plain_csv(src)


def _read_plain_csv(src: SourceSpec):
    header = read_csv_header(src.path, src.delimiter, src.encoding)
    return header, iter_csv_rows(src.path, src.delimiter, src.encoding,
                                 len(header), parse=src.parse)


def _target_label(src: SourceSpec, target: TargetSpec) -> str:
    """多表写出同一个文件时,每个源对应的 Excel Sheet 名 / JSON 一级键。"""
    base = target.target_name(src.table) or Path(src.path).stem or "Sheet1"
    return safe_sheet_title(base) if target.kind == "Excel" else base


def _ensure_fields(header, src: SourceSpec) -> list:
    """空表 / 空集合没有任何字段可写:提前给出中文提示,避免下游拼出非法 SQL。"""
    if not header:
        raise ValueError(f"数据源【{src.label()}】里没有任何字段(可能是空表 / 空集合),无法转换")
    return header


def _write_multi_table(sources, target, conn, adapter, log, verify_sink=None) -> int:
    """多表写出同一个 Excel / JSON 文件(每个源一个 Sheet / 一级键)。"""
    def tables():
        for src in sources:
            log(f"读取数据源: {src.label()}")
            header, rows, closer = open_reader(src, conn=conn, adapter=adapter)
            recorder = NumericRecorder() if target.verify else None
            try:
                header = _ensure_fields(apply_field_mapping(header, target, log), src)
                if recorder is not None:
                    if verify_sink is not None:
                        verify_sink.append((src, header, recorder))
                    rows = recording_rows(rows, header, recorder)
                if target.kind == "Excel":
                    yield _target_label(src, target), header, rows
                else:
                    # MongoDB 来源写文档数组式(直接映射),关系库来源写行记录式
                    yield _target_label(src, target), header, rows, src.kind == "Mongo"
            finally:
                closer()

    writer = write_excel_multi_file if target.kind == "Excel" else write_json_multi_file
    return writer(target.path, tables(), mode=target.mode, log=log)


def apply_field_mapping(header, target: TargetSpec, log=None) -> list:
    """按字段映射重命名输出字段,并保证结果唯一(重名会让写入歧义)。"""
    mapped = target.fields.rename_header(header)
    if len(set(mapped)) != len(mapped):
        duplicates = sorted({name for name in mapped if mapped.count(name) > 1})
        raise ValueError(f"字段映射后出现重名字段: {', '.join(duplicates)}。请调整映射关系")
    if mapped != header and log:
        log(f"  字段映射: {target.fields.describe()}")
    return mapped


def _write_one(src, header, rows, target, conn, adapter, cfg, log, index=1) -> int:
    """把单个输入源写入输出目标,返回写入行数(Writer 统一接口)。

    ``header`` 应是已应用字段映射后的目标表头(由 :func:`convert` 统一处理)。
    """
    if target.kind == "Mongo":
        collection = target.table or target.target_name(src.table)
        if not collection:
            raise ValueError("未指定目标集合名")
        written = write_mongo_collection(conn, adapter, collection, header, rows,
                                         src.mode or target.mode, target.batch, log)
        log(f"✅ 集合【{collection}】写入 {written:,} 行")
        return written

    if target.kind == "SQL":
        table = target.table or target.target_name(src.table)
        if not table:
            raise ValueError("未指定目标表名")
        mode = src.mode or target.mode
        if src.kind == "SQL":
            # 源游标与目标 DDL 不能共用一条连接(PostgreSQL 命名游标会与 DDL 冲突)
            sample, chained = take_sample(rows, SAMPLE_SIZE, len(header))
            out_conn = adapter.connect(cfg)
            try:
                written = write_sql_table(out_conn, adapter, table, header, chained,
                                          mode, target.batch, log, sample=sample)
            finally:
                _close_quietly(out_conn)
        else:
            written = write_sql_table(conn, adapter, table, header, rows,
                                      mode, target.batch, log)
        log(f"✅ 表【{table}】写入 {written:,} 行(模式: {mode})")
        return written

    if target.kind == "CSV" and target.directory:
        path = csv_output_path(target.path, src, index, target)
    else:
        path = target.path

    if target.kind == "CSV":
        written = write_csv_file(path, header, rows, delimiter=target.delimiter,
                                 encoding=target.encoding, mode=target.mode, log=log)
        log(f"✅ CSV 导出完成: {path}(共 {written:,} 行)")
        return written
    if target.kind == "Excel":
        sheet = safe_sheet_title(target.sheet or target.target_name(src.table)
                                 or Path(src.path).stem)
        written = write_excel_file(path, sheet, header, rows, mode=target.mode, log=log)
        log(f"✅ Excel 导出完成: {path}(Sheet [{sheet}],共 {written:,} 行)")
        return written
    if target.kind == "JSON":
        key = (target.key or target.sheet or target.target_name(src.table)
               or Path(src.path).stem or "data")
        written = write_json_file(path, key, header, rows, mode=target.mode, log=log,
                                  documents=src.kind == "Mongo")
        log(f"✅ JSON 导出完成: {path}(键 [{key}],共 {written:,} 行)")
        return written
    raise ValueError(f"不支持的输出类型: {target.kind!r}")


def recording_rows(rows, header, recorder: NumericRecorder):
    """包装行生成器,顺手记录数值字段的原始字面量(按行有界采样)。"""
    for row in rows:
        recorder.record_row(header, row)
        yield row


def readback_source(src, target: TargetSpec, header, index: int = 1,
                    multi: bool = False) -> SourceSpec:
    """构造“回读输出”用的输入描述。

    ``multi=True`` 表示本次是把多个源写进同一个 Excel / JSON 文件(每个源一个
    Sheet / 一级键),名称必须与 :func:`_write_multi_table` 的规则一致;否则会把
    本次写出的 Sheet 找错,回读校验被整段跳过。
    """
    if target.kind == "Mongo":
        return SourceSpec(kind="Mongo", table=target.table or target.target_name(src.table),
                          columns=list(header))
    if target.kind == "SQL":
        return SourceSpec(kind="SQL", table=target.table or target.target_name(src.table),
                          columns=list(header))
    if target.kind == "CSV":
        path = csv_output_path(target.path, src, index, target) if target.directory else target.path
        return SourceSpec(kind="CSV", path=path, delimiter=target.delimiter,
                          encoding=target.encoding)
    if target.kind == "Excel":
        sheet = (_target_label(src, target) if multi else
                 safe_sheet_title(target.sheet or target.target_name(src.table)
                                  or Path(src.path).stem))
        return SourceSpec(kind="Excel", path=target.path, table=sheet)
    if target.kind == "JSON":
        key = (_target_label(src, target) if multi else
               (target.key or target.sheet or target.target_name(src.table)
                or Path(src.path).stem or "data"))
        return SourceSpec(kind="JSON", path=target.path, table=key)
    raise ValueError(f"不支持的输出类型: {target.kind!r}")


# 追加/合并写入已有目标时的回读扫描上限(找齐样本会提前结束,不会真的扫满)
READBACK_SCAN_LIMIT = 200_000


def verify_readback(src, target: TargetSpec, source_header, recorder: NumericRecorder,
                    conn=None, adapter=None, index: int = 1, multi: bool = False,
                    limit: int = DEFAULT_SAMPLE_LIMIT, log=log_noop) -> FidelityReport:
    """回读输出并做数值等值 / 文本等值比对,返回校验报告。

    目标里可能**已有旧数据**(追加写入已有数据库表 / 合并写入已有文件),本次写的行
    不一定在开头;只回读前 ``limit`` 行会把本次的数据误判成"(缺失)"。因此非 replace
    模式下按更大的扫描窗口回读、找齐样本后提前结束,做与位置无关的多重集比对。
    """
    report = FidelityReport()
    if not recorder.active:
        return report
    spec = readback_source(src, target, source_header, index, multi=multi)
    try:
        _, rows, closer = open_reader(spec, conn=conn, adapter=adapter)
    except Exception as exc:                     # 回读失败不应掩盖写入结果
        report.add_note(f"回读输出失败,已跳过数值校验: {exc}")
        return report
    appending = (src.mode or target.mode) != "replace"
    window = READBACK_SCAN_LIMIT if appending else limit
    try:
        collected = {}
        scanned = 0
        for row in rows:
            if scanned >= window:
                break
            scanned += 1
            for name, value in zip(source_header, row):
                collected.setdefault(str(name), []).append(value)
            if appending and scanned % 500 == 0 and _samples_covered(collected, recorder):
                break
    finally:
        closer()
    if appending and scanned >= window:
        report.add_note(f"目标里已有旧数据,回读只扫描了前 {window:,} 行;"
                        f"未出现的样本按缺失计")
    recorder.compare_with(collected, report)
    for line in report.summary_lines():
        log(line, "ok" if report.ok else "err")
    return report


def _samples_covered(collected: dict, recorder: NumericRecorder) -> bool:
    """回读扫描过程中判断:所有采样值是否都已在回读结果里出现(可提前结束扫描)。"""
    for name, originals in recorder.samples.items():
        if not originals:
            continue
        seen = Counter(str(value) for value in collected.get(name, []) if value is not None)
        if Counter(originals) - seen:
            return False
    return True


def convert(sources, target: TargetSpec, conn=None, adapter=None,
            cfg: DBConfig | None = None, log=log_noop) -> RunResult:
    """把 sources 中的每个输入源按 target 的设定写出(界面层与命令行层共用)。

    * 输入涉及数据库时必须提供 conn / adapter;
    * 输出为数据库时必须提供 adapter,cfg 用于另开一条输出连接(避免源游标与 DDL 冲突);
    * target.verify=True 时写入后回读输出,做长浮点数保真校验,未通过则抛
      :class:`dt_numeric.FidelityError`(可用 target.allow_mismatch 放行)。
    """
    sources = list(sources)
    if not sources:
        raise ValueError("没有可转换的输入源")
    needs_db_in = any(src.kind in DB_KINDS for src in sources)
    needs_db_out = target.kind in DB_KINDS
    if (needs_db_in or needs_db_out) and (adapter is None or conn is None):
        raise ValueError("涉及数据库的转换需要先建立数据库连接")

    pending, skipped = [], []
    for src in sources:
        (skipped if (src.mode or target.mode) == "skip" else pending).append(src)
    for src in skipped:
        src.status = "skipped"
        log(f"⏭ 已跳过: {src.label()}")
    result = RunResult(skipped=len(skipped))
    if not pending:
        raise ValueError("所有数据源都被跳过,没有可执行的转换")

    if needs_db_out and target.table and len(pending) > 1:
        raise ValueError("多表拷贝时目标表名必须留空(将按源表同名建表)")
    if target.kind == "CSV" and not target.directory and len(pending) > 1:
        raise ValueError("多个数据源写入 CSV 时请使用目录模式(每个表一个 .csv 文件),"
                         "否则后写入的表会覆盖前一个")
    if target.directory:
        os.makedirs(target.path, exist_ok=True)
        result.outputs.append(target.path)
    elif target.path:
        Path(target.path).parent.mkdir(parents=True, exist_ok=True)
        result.outputs.append(target.path)

    if target.kind in ("Excel", "JSON") and len(pending) > 1:
        verify_sink = []
        result.rows = _write_multi_table(pending, target, conn, adapter, log, verify_sink)
        result.sources = len(pending)
        for src in pending:
            src.status = "done"
        for src, header, recorder in verify_sink:
            if recorder.active:
                result.fidelity.merge(verify_readback(src, target, header, recorder,
                                                      conn=conn, adapter=adapter,
                                                      multi=True, log=log))
        if not result.fidelity.ok and not target.allow_mismatch:
            raise FidelityError(result.fidelity)
        return result

    for index, src in enumerate(pending, start=1):
        log(f"读取数据源: {src.label()}")
        header, rows, closer = open_reader(src, conn=conn, adapter=adapter)
        recorder = NumericRecorder() if target.verify else None
        try:
            header = _ensure_fields(apply_field_mapping(header, target, log), src)
            if recorder is not None:
                rows = recording_rows(rows, header, recorder)
            src.rows = _write_one(src, header, rows, target, conn, adapter, cfg, log, index)
        finally:
            closer()
        if recorder is not None and recorder.active:
            log("  正在回读输出并校验数值保真 ...")
            result.fidelity.merge(verify_readback(src, target, header, recorder,
                                                  conn=conn, adapter=adapter,
                                                  index=index, log=log))
        src.status = "done"
        result.rows += src.rows
        result.sources += 1

    if not result.fidelity.ok and not target.allow_mismatch:
        raise FidelityError(result.fidelity)
    return result


def run_self_test(log_path=None) -> bool:
    """无界面自检:验证依赖、核心算法与文件互转管道,结果写入 selftest.log。"""
    log_path = Path(log_path) if log_path else Path.cwd() / "selftest.log"
    lines = []

    for module_name in ("pymysql", "psycopg2", "openpyxl", "paramiko", "pymongo"):
        try:
            __import__(module_name)
            lines.append(f"[OK]   import {module_name}")
        except Exception as exc:
            level = "SKIP" if module_name in ("paramiko", "pymongo") else "FAIL"
            lines.append(f"[{level}] import {module_name}: {exc}")

    # 表头解析(含重名/空列/尾部空列的边界)
    try:
        cases = {("A", "A", "A"): ["A", "A_2", "A_3"],
                 ("A", "A_2", "A"): ["A", "A_2", "A_3"],
                 ("A", None, "C"): ["A", "column_2", "C"],
                 ("A", None, ""): ["A"],
                 ("A", "B", None): ["A", "B"]}
        ok = all(parse_header(list(raw)) == want for raw, want in cases.items())
        lines.append(f"[{'OK' if ok else 'FAIL'}]   parse_header 去重/补名/去尾空列")
    except Exception as exc:
        lines.append(f"[FAIL] parse_header: {exc}")

    work = log_path.parent / "_selftest"
    shutil.rmtree(work, ignore_errors=True)

    try:
        work.mkdir(parents=True, exist_ok=True)

        # Excel 写读往返(覆盖 openpyxl + jsonb 复合值 + 非法字符)
        tmp = work / "roundtrip.xlsx"
        wb = Workbook(write_only=True)
        write_rows_to_worksheet(
            wb.create_sheet(title="自检"), ["id", "name", "device"],
            iter([[1, "张三", {"ip": "192.168.1.59", "os": "XiaomiNote8"}],
                  [2, "李\u0000四", [1, 2, 3]]]))
        wb.save(tmp)
        _close_quietly(wb)
        wb2 = load_workbook(tmp, read_only=True, data_only=True)
        rows = list(wb2[wb2.sheetnames[0]].iter_rows(values_only=True))
        _close_quietly(wb2)
        ok = (len(rows) == 3 and rows[1][0] == 1 and "192.168.1.59" in str(rows[1][2])
              and rows[2][1] == "李四")
        lines.append(f"[{'OK' if ok else 'FAIL'}]   Excel 写读往返(含 jsonb/控制字符): {len(rows)} 行")

        # JSON 写读/合并往返 + 旧版列数组兼容
        tmp = work / "roundtrip.json"
        write_json_file(str(tmp), "info", ["Id", "Name"], [[1, "张三"], [2, "李四"]])
        with open(tmp, "r", encoding="utf-8") as fp:
            doc = json.load(fp)
        header, gen = load_json_stream(str(tmp), "info")
        got = list(gen)
        ok = (doc["info"]["HeaderFields"] == ["Id", "Name"] and got[0] == [1, "张三"]
              and header == ["Id", "Name"])
        write_json_file(str(tmp), "info", ["Id", "Name"], [[3, "王五"]], mode="merge")
        write_json_file(str(tmp), "other", ["X"], [[9]], mode="merge")
        with open(tmp, "r", encoding="utf-8") as fp:
            doc = json.load(fp)
        ok = ok and doc["info"]["Data"]["Row1"] == [3, "王五"] and "other" in doc
        with open(tmp, "w", encoding="utf-8") as fp:
            json.dump({"legacy": {"Id": [7, 8], "Name": ["甲", "乙"]}}, fp, ensure_ascii=False)
        header, gen = load_json_stream(str(tmp), "legacy")
        ok = ok and header == ["Id", "Name"] and list(gen)[1] == [8, "乙"]
        lines.append(f"[{'OK' if ok else 'FAIL'}]   JSON 写读/合并/旧结构兼容")

        # CSV 写读往返(含类型解析与复合值)
        tmp = work / "roundtrip.csv"
        write_csv_file(str(tmp), ["Id", "Name", "Device"],
                       iter([[1, "张三", {"ip": "192.168.1.59"}], [2, "李四", [4, 5]]]))
        header = read_csv_header(str(tmp), ",", "utf-8-sig")
        rows = list(iter_csv_rows(str(tmp), ",", "utf-8-sig", len(header), parse=True))
        ok = (header == ["Id", "Name", "Device"] and rows[0][0] == 1
              and rows[1][2] == "[4, 5]")
        lines.append(f"[{'OK' if ok else 'FAIL'}]   CSV 写读往返: {len(rows)} 行")

        # 旧版分节 CSV 读取(向后兼容)
        tmp = work / "sections.csv"
        with open(tmp, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow(["[Sheet:info]"])
            writer.writerow(["Id", "Name"])
            writer.writerow(["1", "张三"])
        header, gen = load_csv_section(str(tmp), "info", ",", "utf-8-sig")
        ok = (scan_csv_sections(str(tmp), ",", "utf-8-sig") == ["info"]
              and header == ["Id", "Name"] and list(gen)[0] == ["1", "张三"])
        lines.append(f"[{'OK' if ok else 'FAIL'}]   分节CSV读取(兼容旧文件)")

        # 端到端转换管道:CSV → JSON → Excel → CSV(目录模式)→ CSV(字段过滤)
        src_csv = work / "pipeline.csv"
        write_csv_file(str(src_csv), ["Id", "Name", "Device"],
                       iter([[1, "张三", "phone"], [2, "李四", "pad"]]))
        out_json = work / "pipeline.json"
        convert([SourceSpec(kind="CSV", path=str(src_csv), parse=True)],
                TargetSpec(kind="JSON", path=str(out_json), key="info"))
        out_xlsx = work / "pipeline.xlsx"
        convert([SourceSpec(kind="JSON", path=str(out_json), table="info")],
                TargetSpec(kind="Excel", path=str(out_xlsx), sheet="info"))
        out_dir = work / "csv_dir"
        convert([SourceSpec(kind="Excel", path=str(out_xlsx), table="info",
                            columns=["Id", "Name"])],
                TargetSpec(kind="CSV", path=str(out_dir), directory=True))
        produced = out_dir / "info.csv"
        header = read_csv_header(str(produced), ",", "utf-8-sig")
        rows = list(iter_csv_rows(str(produced), ",", "utf-8-sig", len(header), parse=True))
        ok = (produced.exists() and header == ["Id", "Name"] and rows[0] == [1, "张三"]
              and json_top_keys(str(out_json)) == ["info"]
              and "info" in workbook_sheet_headers(str(out_xlsx)))
        lines.append(f"[{'OK' if ok else 'FAIL'}]   转换管道 CSV→JSON→Excel→CSV目录(含字段过滤)")
    except Exception as exc:
        lines.append(f"[FAIL] 转换管道: {exc}")
    finally:
        shutil.rmtree(work, ignore_errors=True)

    # 长浮点数保真:CSV → JSON → CSV,文本必须逐字一致且不出现科学计数法
    numeric_work = log_path.parent / "_selftest_numeric"
    try:
        shutil.rmtree(numeric_work, ignore_errors=True)
        numeric_work.mkdir(parents=True, exist_ok=True)
        samples = ["12345678901234567890.12345", "1.500", "9007199254740993", "1E+20"]
        expected = ["12345678901234567890.12345", "1.500", "9007199254740993",
                    "100000000000000000000"]
        numeric_csv = numeric_work / "numeric.csv"
        numeric_json = numeric_work / "numeric.json"
        numeric_back = numeric_work / "numeric_back.csv"
        write_csv_file(str(numeric_csv), ["Amount"], iter([[value] for value in samples]))
        convert([SourceSpec(kind="CSV", path=str(numeric_csv), parse=True)],
                TargetSpec(kind="JSON", path=str(numeric_json), key="info"))
        convert([SourceSpec(kind="JSON", path=str(numeric_json), table="info")],
                TargetSpec(kind="CSV", path=str(numeric_back)))
        text = numeric_json.read_text(encoding="utf-8")
        got = [line.split(",")[0]
               for line in numeric_back.read_text(encoding="utf-8-sig").strip().splitlines()[1:]]
        ok = got == expected and "e+" not in text and "E+" not in text
        lines.append(f"[{'OK' if ok else 'FAIL'}]   长浮点数保真(CSV→JSON→CSV 文本一致,无科学计数法)")
    except Exception as exc:
        lines.append(f"[FAIL] 长浮点数保真: {exc}")
    finally:
        shutil.rmtree(numeric_work, ignore_errors=True)

    # SSH 参数校验(不联网)
    try:
        valid = SSHConfig(enabled=True, host="h", user="u", password="p").validated()
        ok = valid.port == SSH_DEFAULT_PORT and valid.host == "h"
        try:
            SSHConfig(enabled=True, host="", user="").validated()
            ok = False                        # 缺少主机名时必须报错
        except ValueError:
            pass
        lines.append(f"[{'OK' if ok else 'FAIL'}]   SSH 参数校验"
                     f"(paramiko {'已安装' if paramiko_available() else '未安装'})")
    except Exception as exc:
        lines.append(f"[FAIL] SSH 参数校验: {exc}")

    # 若提供数据库连接环境变量,则做真实连接自检
    for env_prefix, adapter in (("SELFTEST_MYSQL", ADAPTERS["MySQL"]),
                                ("SELFTEST_PG", ADAPTERS["PostgreSQL"])):
        cfg = DBConfig.from_mapping({
            "host": os.environ.get(f"{env_prefix}_HOST", "127.0.0.1"),
            "port": os.environ.get(f"{env_prefix}_PORT") or adapter.DEFAULT_PORT,
            "user": os.environ.get(f"{env_prefix}_USER", ""),
            "password": os.environ.get(f"{env_prefix}_PASSWORD", ""),
            "database": os.environ.get(f"{env_prefix}_DATABASE", ""),
        })
        if not cfg.user or not cfg.database:
            lines.append(f"[SKIP] {adapter.NAME} 连接自检(未设置 {env_prefix}_* 环境变量)")
            continue
        try:
            conn = adapter.connect(cfg.validated(adapter))
            with conn.cursor() as cur:
                cur.execute("SELECT 1")
                cur.fetchone()
            conn.close()
            lines.append(f"[OK]   {adapter.NAME} 连接 SELECT 1")
        except Exception as exc:
            lines.append(f"[FAIL] {adapter.NAME} 连接: {exc}")

    # MongoDB 连接自检(设置 SELFTEST_MONGO_* 环境变量时)
    if not pymongo_available():
        lines.append("[SKIP] MongoDB 连接自检(未安装 pymongo)")
    elif not (mongo_db := os.environ.get("SELFTEST_MONGO_DATABASE")):
        lines.append("[SKIP] MongoDB 连接自检(未设置 SELFTEST_MONGO_* 环境变量)")
    else:
        try:
            mongo_cfg = MongoConfig.from_mapping({
                "host": os.environ.get("SELFTEST_MONGO_HOST", "127.0.0.1"),
                "port": os.environ.get("SELFTEST_MONGO_PORT") or 27017,
                "user": os.environ.get("SELFTEST_MONGO_USER", ""),
                "password": os.environ.get("SELFTEST_MONGO_PASSWORD", ""),
                "database": mongo_db,
                "uri": os.environ.get("SELFTEST_MONGO_URI", ""),
                "timeout": 10,
            })
            mongo_adapter = ADAPTERS["MongoDB"]
            database = mongo_adapter.connect(mongo_cfg)
            try:
                lines.append(f"[OK]   {mongo_adapter.server_version(database)} 连接自检"
                             f"({len(mongo_adapter.list_tables(database))} 个集合)")
            finally:
                mongo_adapter.close(database)
        except Exception as exc:
            lines.append(f"[FAIL] MongoDB 连接: {exc}")

    log_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return all(not line.startswith("[FAIL]") for line in lines)


if __name__ == "__main__":  # pragma: no cover - 直接运行本文件等价于自检
    raise SystemExit(0 if run_self_test() else 1)
